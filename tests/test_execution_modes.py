"""Modos de execução (rápido × completo), limite de requisições e chamadas mínimas."""

from __future__ import annotations

import json
from dataclasses import replace
from datetime import datetime
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import load_workbook

from src.api_client import CarteiraAdvisorClient
from src.cadastral import CadastralStatus
from src.config import MoneyUnit, Settings, load_settings
from src.errors import ExitCode, IntegrationError
from src.estimate import ZeroCategory
from src.exporter import HEADERS, SHEET_NAME
from src.pipeline import run
from src.reconciliation import AuditStatus
from src.summary import render_summary
from tests.conftest import FakeApi, make_client, make_cnpj, make_pl_entry, make_wallet_detail

NOW = datetime(2026, 10, 28, 14, 0, 0)


def fast(settings: Settings) -> Settings:
    return replace(settings, detailed_audit=False)


def run_with(settings: Settings, api: Any):
    return run(settings, now=NOW, client_factory=lambda _: api)


# --- Modo rápido -------------------------------------------------------------


def test_detailed_audit_is_disabled_by_default() -> None:
    assert load_settings(env={}).detailed_audit is False
    assert load_settings(env={"CARTEIRA_ADVISOR_DETAILED_AUDIT": "true"}).detailed_audit is True
    assert load_settings(env={}).max_requests_per_minute == 50


def test_fast_mode_makes_only_the_necessary_calls(settings: Settings) -> None:
    entries = [make_pl_entry(1, 10), make_pl_entry(2, 5)]
    entries[1]["wallets"][0]["walletInfoId"] = 2000
    api = FakeApi([make_client(1), make_client(2, cpfCnpj=make_cnpj(2))], entries, total=15)
    result = run_with(fast(settings), api)

    assert api.calls == ["clients", "consultants", "pl"]  # sem detalhes nem segunda leitura
    assert result.exit_code is ExitCode.SUCCESS
    assert result.detailed_audit is False
    assert result.fx is None and result.stability is None and result.diagnostic_path is None
    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED
    assert result.cadastral.status is CadastralStatus.APPROVED  # IDs conferidos com uma leitura
    sheet = load_workbook(result.output_path)[SHEET_NAME]
    assert tuple(c.value for c in sheet[1]) == HEADERS and len(HEADERS) == 8
    text = render_summary(result, MoneyUnit.REAIS)
    assert "AUDITORIA FINANCEIRA DETALHADA: NÃO EXECUTADA (CARTEIRA_ADVISOR_DETAILED_AUDIT=false)" in text
    assert "VERIFICAÇÃO INDEPENDENTE" not in text and "VERIFICAÇÃO TEMPORAL" not in text


def test_fast_mode_zero_patrimony_is_listed_but_not_evaluated(settings: Settings) -> None:
    api = FakeApi([make_client(1, name="Zero Ficticio", currentValue=0)], [make_pl_entry(1, 0)], total=0)
    result = run_with(fast(settings), api)
    assert [z.category for z in result.zero_patrimony] == [ZeroCategory.NOT_EVALUATED]
    text = render_summary(result, MoneyUnit.REAIS)
    assert "PF ativas com PL zero (informativo; pode ser intencional): 1" in text
    assert "Zero Ficticio (ID 1): PL zero (consolidação não avaliada" in text
    line = next(x for x in text.splitlines() if "Zero Ficticio (ID 1)" in x)
    assert "1 carteira(s); cadastro R$ 0,00" in line
    assert "SEM_MOEDA" not in line and "nunca" not in line  # nada afirmado sem consulta


def test_fast_mode_survives_financial_endpoint_failure(settings: Settings) -> None:
    api = FakeApi([make_client(1)], [], total=0, fail_on="pl", error=IntegrationError("HTTP 503"))
    result = run_with(fast(settings), api)
    assert result.exit_code is ExitCode.SUCCESS
    assert result.cadastral.status is CadastralStatus.NOT_VERIFIED
    assert (settings.output_dir / "Aniversariantes_2026-10-28_140000.xlsx").exists()
    assert "Consolidado indisponível: HTTP 503" in render_summary(result, MoneyUnit.REAIS)


def test_full_mode_calls(settings: Settings) -> None:
    api = FakeApi([make_client(1)], [make_pl_entry(1, 10)], total=10)
    run_with(settings, api)  # fixture: detailed_audit=True
    assert api.calls == ["clients", "consultants", "pl", "details", "pl"]


def test_patrimony_is_recomputed_on_every_run(settings: Settings) -> None:
    first = run_with(
        fast(settings), FakeApi([make_client(1, currentValue=100)], [make_pl_entry(1, 100)], total=100)
    )
    later = replace(settings, output_dir=settings.output_dir / "segunda")
    second = run_with(
        fast(later), FakeApi([make_client(1, currentValue=250)], [make_pl_entry(1, 250)], total=250)
    )
    rows = [
        list(load_workbook(r.output_path)[SHEET_NAME].iter_rows(min_row=2, values_only=True))
        for r in (first, second)
    ]
    assert rows[0][0][2] == 100 and rows[1][0][2] == 250


# --- Cliente HTTP real com sessão simulada: contagem exata de requisições ------


class Response:
    def __init__(self, body: Any, status: int = 200, headers: dict | None = None) -> None:
        self.status_code = status
        self.headers = headers or {}
        self._text = json.dumps(body, default=float)

    def json(self, **kwargs: Any) -> Any:
        return json.loads(self._text, **kwargs)


class Session:
    def __init__(self, clients: list, entries: list) -> None:
        self.headers: dict = {}
        self.urls: list[str] = []
        self.clients, self.entries = clients, entries

    def get(self, url: str, params: Any = None, timeout: float | None = None) -> Response:
        self.urls.append(url)
        if url.endswith("/company-api/clients"):
            return Response(self.clients)
        if url.endswith("/company-api/consultants"):
            return Response([{"id": 1, "name": "Consultora Exemplo"}])
        if url.endswith("/consolidated-pl"):
            page = {"content": self.entries, "totalPages": 1, "totalElements": len(self.entries), "number": 0}
            return Response({"totalCurrentValue": 30, "clients": page})
        wallet_id = int(url.rsplit("/", 1)[1])
        return Response(make_wallet_detail(wallet_id))

    def __getattr__(self, name: str) -> Any:
        if name in {"post", "put", "patch", "delete", "request"}:
            raise AssertionError(f"verbo proibido: {name}")
        raise AttributeError(name)


def real_client_run(settings: Settings, detailed: bool) -> tuple[Any, Session]:
    clients = [make_client(1), make_client(2)]
    entries = [make_pl_entry(1, 10), make_pl_entry(2, 20)]
    entries[1]["wallets"][0]["walletInfoId"] = 2000
    session = Session(clients, entries)
    custom = replace(settings, detailed_audit=detailed)
    result = run(custom, now=NOW, client_factory=lambda s: CarteiraAdvisorClient(s, session=session))
    return result, session


def test_fast_mode_sends_exactly_three_requests(settings: Settings) -> None:
    result, session = real_client_run(settings, detailed=False)
    assert len(session.urls) == 3 == result.api_requests
    assert "Requisições ao Carteira Advisor nesta execução: 3" in render_summary(result, MoneyUnit.REAIS)


def test_full_mode_sends_one_request_per_wallet_plus_two_reads(settings: Settings) -> None:
    result, session = real_client_run(settings, detailed=True)
    # clientes + consultores + consolidado + 2 carteiras + consolidado (estabilidade)
    assert len(session.urls) == 6 == result.api_requests
    assert sum(url.endswith("/consolidated-pl") for url in session.urls) == 2


# --- Limite de 50 requisições por minuto -------------------------------------------


def test_requests_are_spaced_to_respect_the_per_minute_limit(
    settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    clock = {"now": 1000.0}
    sleeps: list[float] = []

    def sleep(seconds: float) -> None:
        sleeps.append(round(seconds, 6))
        clock["now"] += seconds

    monkeypatch.setattr("src.api_client.time.monotonic", lambda: clock["now"])
    monkeypatch.setattr("src.api_client.time.sleep", sleep)
    limited = replace(settings, max_requests_per_minute=50, wallet_detail_interval_seconds=0)
    session = Session([make_client(1)], [make_pl_entry(1, 10)])
    api = CarteiraAdvisorClient(limited, session=session)
    api.fetch_clients()
    api.fetch_consultants()
    api.fetch_consolidated_pl()
    assert sleeps == [1.2, 1.2]  # 60 s / 50 requisições
    assert api.request_count == 3


def test_limit_includes_retries_after_429(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    clock = {"now": 0.0}
    sleeps: list[float] = []

    def sleep(seconds: float) -> None:
        sleeps.append(round(seconds, 6))
        clock["now"] += seconds

    monkeypatch.setattr("src.api_client.time.monotonic", lambda: clock["now"])
    monkeypatch.setattr("src.api_client.time.sleep", sleep)
    responses = iter([Response({}, 429, {"Retry-After": "0.5"}), Response([make_client(1)])])

    class Flaky(Session):
        def get(self, url: str, params: Any = None, timeout: float | None = None) -> Response:
            self.urls.append(url)
            return next(responses)

    api = CarteiraAdvisorClient(
        replace(settings, max_requests_per_minute=50, max_retries=2), session=Flaky([], [])
    )
    assert len(api.fetch_clients()) == 1
    # Retry-After 0,5 s; depois o limitador completa o espaçamento até 1,2 s.
    assert sleeps == [0.5, 0.7]
    assert api.request_count == 2


def test_full_audit_respects_limit_end_to_end(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    clock = {"now": 0.0}
    sent_at: list[float] = []

    monkeypatch.setattr("src.api_client.time.monotonic", lambda: clock["now"])
    monkeypatch.setattr("src.api_client.time.sleep", lambda s: clock.__setitem__("now", clock["now"] + s))
    clients = [make_client(i) for i in range(1, 61)]
    entries = [make_pl_entry(i, 10) for i in range(1, 61)]
    for index, entry in enumerate(entries):
        entry["wallets"][0]["walletInfoId"] = 5000 + index

    class Timed(Session):
        def get(self, url: str, params: Any = None, timeout: float | None = None) -> Response:
            sent_at.append(clock["now"])
            return super().get(url, params, timeout)

    session = Timed(clients, entries)
    custom = replace(
        settings, detailed_audit=True, max_requests_per_minute=50, wallet_detail_interval_seconds=0
    )
    run(custom, now=NOW, client_factory=lambda s: CarteiraAdvisorClient(s, session=session))
    assert len(sent_at) == 64  # 3 + 60 carteiras + segunda leitura
    for index in range(len(sent_at) - 50):
        assert sent_at[index + 50] - sent_at[index] >= 60 - 1e-9  # nunca mais de 50 em 60 s


def test_total_value_stays_dynamic_and_not_cached(settings: Settings) -> None:
    from src import pipeline

    assert not [name for name in dir(pipeline) if "cache" in name.lower() and not name.startswith("__")]
    result = run_with(fast(settings), FakeApi([make_client(1)], [make_pl_entry(1, 10)], total=Decimal("10")))
    assert result.reconciliation.company_total == Decimal("10")

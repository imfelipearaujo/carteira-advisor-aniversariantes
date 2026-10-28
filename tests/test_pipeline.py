"""End-to-end flow with a fake API.

The operational PF spreadsheet is published regardless of the financial
audit; only configuration, cadastral collection and export failures block it.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

import main
from src.cadastral import ActivityStatus, CadastralStatus
from src.config import MoneyUnit, Settings
from src.errors import CollectionIntegrityError, ConfigurationError, ExitCode, IntegrationError
from src.exporter import SHEET_NAME
from src.pipeline import run
from src.reconciliation import AuditStatus
from src.summary import format_brl, render_summary
from src.transformer import PersonType
from tests.conftest import FakeApi, make_client, make_cnpj, make_pl_entry

NOW = datetime(2026, 10, 28, 14, 0, 0)
EXPECTED_FILE = "Aniversariantes_2026-10-28_140000.xlsx"


def published_files(settings: Settings) -> list[Path]:
    if not settings.output_dir.exists():
        return []
    return sorted(p for p in settings.output_dir.rglob("*") if p.is_file())


def run_with(settings: Settings, api: FakeApi):
    return run(settings, now=NOW, client_factory=lambda _: api)


def sheet_rows(path: Path) -> list[tuple]:
    return list(load_workbook(path)[SHEET_NAME].iter_rows(min_row=2, values_only=True))


# --- Financial audit: still computed, never blocks the spreadsheet ----------


def test_scenario_a_publishes_report_and_audit_is_approved(settings: Settings) -> None:
    api = FakeApi(
        [make_client(1, birthdate="1990-03-01"), make_client(2, birthdate=None)],
        [make_pl_entry(1, 1000), make_pl_entry(2, 1000)],
        total=2000,
    )
    result = run_with(settings, api)

    assert result.exit_code is ExitCode.SUCCESS
    assert result.reconciliation.status is AuditStatus.APPROVED
    assert result.output_path == settings.output_dir / EXPECTED_FILE
    assert published_files(settings) == [result.output_path]
    sheet = load_workbook(result.output_path)[SHEET_NAME]
    assert sheet.max_row == 3
    assert sheet.cell(3, 2).value == "PENDENTE"
    assert (result.clients_found, result.clients_exported, result.valid_birthdays) == (2, 2, 1)


def test_scenario_b_divergence_still_publishes_operational_report(settings: Settings) -> None:
    api = FakeApi([make_client(1), make_client(2)], [make_pl_entry(1, 1000), make_pl_entry(2, 1000)], 5000)
    result = run_with(settings, api)

    assert result.reconciliation.status is AuditStatus.DIVERGENCE
    assert result.reconciliation.difference == Decimal(3000)
    assert result.exit_code is ExitCode.SUCCESS
    assert published_files(settings) == [settings.output_dir / EXPECTED_FILE]


def test_scenario_c_overlap_is_not_approved_and_sheet_uses_current_value(
    exclusive_settings: Settings,
) -> None:
    api = FakeApi([make_client(1, currentValue=1234)], [make_pl_entry(1, 1000, 2000, 3000)], total=6000)
    result = run_with(exclusive_settings, api)

    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED
    (row,) = sheet_rows(result.output_path)
    # Neither 6000 (naive wallet sum) nor blank: the operational source is currentValue.
    assert row[2] == 1234
    assert row[7] == "OK"


def test_multiple_wallets_unconfirmed_policy_not_approved(settings: Settings) -> None:
    api = FakeApi([make_client(1)], [make_pl_entry(1, 1000, 2500)], total=3500)
    result = run_with(settings, api)
    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED
    assert result.output_path is not None


def test_multiple_wallets_exclusive_policy_approved(exclusive_settings: Settings) -> None:
    api = FakeApi([make_client(1)], [make_pl_entry(1, 1000, 2500)], total=3500)
    result = run_with(exclusive_settings, api)
    assert result.reconciliation.status is AuditStatus.APPROVED


def test_client_absent_from_patrimony_is_kept_but_not_approved(settings: Settings) -> None:
    api = FakeApi([make_client(1), make_client(2)], [make_pl_entry(1, 1000)], total=1000)
    result = run_with(settings, api)
    assert result.clients_exported == 2
    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED


def test_orphan_patrimony_record_blocks_approval(settings: Settings) -> None:
    api = FakeApi([make_client(1)], [make_pl_entry(1, 1000), make_pl_entry(77, 0)], total=1000)
    result = run_with(settings, api)
    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED
    assert result.reconciliation.orphan_patrimony_ids == [77]


def test_sheet_uses_current_value_and_audit_uses_wallets(settings: Settings) -> None:
    api = FakeApi([make_client(1, currentValue=9999)], [make_pl_entry(1, 1000)], total=1000)
    result = run_with(settings, api)
    # The audit keeps its own wallet-based figure; the sheet shows currentValue.
    assert result.reconciliation.status is AuditStatus.APPROVED
    assert result.reconciliation.clients_total == Decimal(1000)
    assert sheet_rows(result.output_path)[0][2] == 9999


def test_patrimony_endpoint_failure_does_not_block_spreadsheet(settings: Settings) -> None:
    api = FakeApi(
        [make_client(1)],
        [make_pl_entry(1, 1000)],
        total=1000,
        fail_on="pl",
        error=CollectionIntegrityError("página faltante"),
    )
    result = run_with(settings, api)

    assert result.exit_code is ExitCode.SUCCESS
    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED
    assert result.audit_error == "página faltante"
    assert published_files(settings) == [settings.output_dir / EXPECTED_FILE]
    assert "Patrimônio consolidado indisponível" in render_summary(result, MoneyUnit.REAIS)


# --- Blocking failures --------------------------------------------------------


def test_missing_token_makes_no_request_and_no_file(settings: Settings) -> None:
    from dataclasses import replace

    def factory(_: Settings) -> FakeApi:
        pytest.fail("cliente HTTP não deveria ser criado sem token")

    with pytest.raises(ConfigurationError):
        run(replace(settings, token=""), now=NOW, client_factory=factory)
    assert published_files(settings) == []


@pytest.mark.parametrize("stage", ["clients", "consultants"])
def test_cadastral_collection_failure_publishes_nothing(settings: Settings, stage: str) -> None:
    api = FakeApi(
        [make_client(1)],
        [make_pl_entry(1, 1000)],
        total=1000,
        fail_on=stage,
        error=CollectionIntegrityError("página faltante"),
    )
    with pytest.raises(IntegrationError):
        run_with(settings, api)
    assert published_files(settings) == []


def test_export_failure_leaves_no_partial_file(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    from src.errors import ExportError

    def broken_validation(*_: object) -> None:
        raise ExportError("falha simulada")

    monkeypatch.setattr("src.pipeline.validate_excel", broken_validation)
    api = FakeApi([make_client(1)], [make_pl_entry(1, 1000)], total=1000)
    with pytest.raises(ExportError):
        run_with(settings, api)
    assert published_files(settings) == []


# --- PF / PJ classification -----------------------------------------------------


def test_only_individuals_are_exported(settings: Settings) -> None:
    clients = [
        make_client(1, name="Pessoa Fisica Um"),
        make_client(2, name="Empresa Ficticia Ltda", cpfCnpj=make_cnpj(2)),
        make_client(3, name="Documento Ausente", cpfCnpj=None),
        make_client(4, name="Documento Invalido", cpfCnpj="12345678900"),
        make_client(5, name="Pessoa Fisica Dois", active=False),
    ]
    result = run_with(settings, FakeApi(clients, [], total=0))

    names = {row[0] for row in sheet_rows(result.output_path)}
    # "Pessoa Fisica Dois" is a valid CPF but inactive: excluded since the active filter.
    assert names == {"Pessoa Fisica Um"}
    assert (result.clients_found, result.individuals, result.companies_excluded) == (5, 2, 1)
    assert result.unclassified_ids == (3, 4)
    assert result.clients_exported == 1


def test_pf_count_integrity_on_large_mixed_base(settings: Settings) -> None:
    clients = [make_client(i) for i in range(1, 66)]
    clients += [make_client(i, cpfCnpj=make_cnpj(i)) for i in range(66, 78)]
    clients += [make_client(78, cpfCnpj=None), make_client(79, cpfCnpj=None)]
    result = run_with(settings, FakeApi(clients, [], total=0))

    rows = sheet_rows(result.output_path)
    assert (result.individuals, result.companies_excluded, len(result.unclassified_ids)) == (65, 12, 2)
    assert len(rows) == result.clients_exported == 65


def test_summary_shows_classification_and_provisional_patrimony(settings: Settings) -> None:
    clients = [
        make_client(1, currentValue=Decimal("1000.50")),
        make_client(2, currentValue=0),
        make_client(3, cpfCnpj=make_cnpj(3), currentValue=999_999),
        make_client(4, cpfCnpj=None),
    ]
    result = run_with(settings, FakeApi(clients, [], total=0))
    text = render_summary(result, MoneyUnit.REAIS, settings.output_dir.parent)

    # PJ value is excluded from the diagnostic total.
    assert result.provisional_patrimony_total == Decimal("1000.50")
    assert "CLASSIFICAÇÃO CADASTRAL" in text and "  PJ " in text
    assert "Documentos sem classificação segura, fora da planilha: Cliente Ficticio 4 (ID 4)" in text
    assert "PATRIMÔNIO PROVISÓRIO" in text and "NÃO auditado" in text
    assert "R$ 1.000,50" in text
    assert f"output/{EXPECTED_FILE}" in text


# --- Provisional patrimony in the spreadsheet ----------------------------------


def test_patrimony_column_from_current_value_zero_and_missing(settings: Settings) -> None:
    clients = [
        make_client(1, name="A", birthdate="1980-01-01", currentValue=Decimal("2500.75")),
        make_client(2, name="B", birthdate="1980-01-02", currentValue=0),
        make_client(3, name="C", birthdate="1980-01-03", currentValue=None),
    ]
    result = run_with(settings, FakeApi(clients, [], total=0))
    rows = sheet_rows(result.output_path)

    assert [r[2] for r in rows] == [2500.75, 0, None]
    assert [r[7] for r in rows] == ["OK", "OK", "PENDÊNCIA: PATRIMÔNIO"]
    assert result.pending_patrimony == 1


def test_spreadsheet_has_no_provisional_patrimony_note(settings: Settings) -> None:
    result = run_with(settings, FakeApi([make_client(1)], [], total=0))
    workbook = load_workbook(result.output_path)
    assert workbook[SHEET_NAME].cell(1, 3).comment is None
    assert workbook.properties.description != "Patrimônio PROVISÓRIO"


def test_pending_birthdates_are_last(settings: Settings) -> None:
    clients = [
        make_client(1, name="Sem Data", birthdate=None),
        make_client(2, name="Dezembro", birthdate="1970-12-31"),
        make_client(3, name="Invalida", birthdate="1990-02-30"),
        make_client(4, name="Janeiro", birthdate="2000-01-05"),
    ]
    rows = sheet_rows(run_with(settings, FakeApi(clients, [], total=0)).output_path)
    assert [r[0] for r in rows] == ["Janeiro", "Dezembro", "Invalida", "Sem Data"]
    assert [r[1] for r in rows][2:] == ["PENDENTE", "PENDENTE"]


# --- Summary / formatting / main ----------------------------------------------


def test_summary_text(settings: Settings) -> None:
    api = FakeApi([make_client(1)], [make_pl_entry(1, 1000)], total=1000)
    text = render_summary(run_with(settings, api), MoneyUnit.REAIS, settings.output_dir.parent)
    assert "AUDITORIA PATRIMONIAL: APROVADA" in text
    assert "R$ 1.000,00" in text
    assert f"output/{EXPECTED_FILE}" in text


@pytest.mark.parametrize(
    ("value", "unit", "expected"),
    [
        (Decimal("50000000"), MoneyUnit.REAIS, "R$ 50.000.000,00"),
        (Decimal("-1234.5"), MoneyUnit.REAIS, "-R$ 1.234,50"),
        (Decimal("123456"), MoneyUnit.CENTAVOS, "R$ 1.234,56"),
        (Decimal("0.001"), MoneyUnit.REAIS, "R$ 0,001"),
        (None, MoneyUnit.REAIS, "n/d"),
    ],
)
def test_format_brl(value: Decimal | None, unit: MoneyUnit, expected: str) -> None:
    assert format_brl(value, unit) == expected


def test_main_without_token_exits_with_configuration_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    from dataclasses import replace

    from src import config

    base = config.load_settings(env={})
    monkeypatch.setattr(
        main, "load_settings", lambda: replace(base, log_dir=tmp_path / "logs", output_dir=tmp_path)
    )
    monkeypatch.setattr(main, "PROJECT_ROOT", tmp_path)

    assert main.main() == ExitCode.CONFIGURATION_ERROR
    err = capsys.readouterr().err
    assert "CARTEIRA_ADVISOR_TOKEN" in err
    assert not list(tmp_path.glob("*.xlsx"))


def test_clients_missing_from_consolidated_are_audit_pending_not_zero(exclusive_settings: Settings) -> None:
    # Mirrors production: 79 registered clients, 77 in the consolidated patrimony.
    clients = [make_client(i) for i in range(1, 80)]
    entries = [make_pl_entry(i, 100) for i in range(1, 80) if i not in (40, 41)]
    api = FakeApi(clients, entries, total=7700)

    result = run_with(exclusive_settings, api)

    assert (result.clients_found, result.clients_exported, result.patrimony_clients) == (79, 79, 77)
    assert result.missing_patrimony_ids == (40, 41)
    # Even though 77 x 100 == 7700, the missing clients block audit approval...
    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED
    # ...but the operational sheet is published with their cadastral value.
    names = {"Cliente Ficticio 40", "Cliente Ficticio 41"}
    missing_rows = [r for r in sheet_rows(result.output_path) if r[0] in names]
    assert [r[2] for r in missing_rows] == [1000, 1000]


def test_summary_shows_counts_missing_ids_and_unaudited_raw_sum(settings: Settings) -> None:
    api = FakeApi([make_client(1), make_client(2)], [make_pl_entry(1, 1000, 500)], total=2000)
    result = run_with(settings, api)
    text = render_summary(result, MoneyUnit.REAIS)

    assert result.wallet_count == 2
    assert result.raw_wallets_total == Decimal(1500)
    assert "Carteiras encontradas:" in text
    assert "Clientes sem registro no consolidado: Cliente Ficticio 2 (ID 2)" in text
    assert "NÃO auditada" in text
    assert "R$ 500,00" in text
    assert "AUDITORIA PATRIMONIAL: NÃO VERIFICADA" in text


# --- Multi-currency audit (isolated from the spreadsheet) -----------------------


def test_usd_wallet_keeps_audit_unverified_and_never_mixes_currencies(settings: Settings) -> None:
    from tests.conftest import make_wallet_detail

    entries = [make_pl_entry(1, 1000), make_pl_entry(2, 100)]
    entries[1]["wallets"][0]["walletInfoId"] = 2000
    api = FakeApi(
        [make_client(1), make_client(2)],
        entries,
        total=Decimal("1534.62"),
        wallet_details={2000: make_wallet_detail(2000, walletCoin="DOLLAR")},
    )
    result = run_with(settings, api)
    text = render_summary(result, MoneyUnit.REAIS)

    assert result.exit_code is ExitCode.SUCCESS
    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED
    assert result.reconciliation.difference is None
    assert result.currency.totals == {"REAIS": Decimal(1000), "DOLLAR": Decimal(100)}
    assert "DOLLAR 100,00" in text and "NÃO é cotação oficial" in text
    # Neither the mixed raw sum (1100) nor a mixed per-client total is shown in R$.
    assert "Soma bruta" not in text
    assert "Total auditado:" in text and "n/d (moedas)" in text
    assert "R$ 1.100,00" not in text
    assert published_files(settings) == [settings.output_dir / EXPECTED_FILE]


def test_wallet_detail_failure_does_not_block_spreadsheet(settings: Settings) -> None:
    api = FakeApi(
        [make_client(1)],
        [make_pl_entry(1, 1000)],
        total=1000,
        fail_on="details",
        error=IntegrationError("GET /company-api/wallet/1000: HTTP 429 inesperado."),
    )
    result = run_with(settings, api)
    assert result.exit_code is ExitCode.SUCCESS
    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED
    assert any("Moeda das carteiras não verificada" in r for r in result.reconciliation.reasons)
    assert published_files(settings) == [settings.output_dir / EXPECTED_FILE]


def test_disabled_wallet_details_never_approves(settings: Settings) -> None:
    from dataclasses import replace

    api = FakeApi([make_client(1)], [make_pl_entry(1, 1000)], total=1000)
    result = run_with(replace(settings, wallet_details_enabled=False), api)
    assert "details" not in api.calls
    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED
    assert result.currency is None


def test_unexpected_audit_error_still_publishes(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    def boom(*_: object) -> None:
        raise RuntimeError("bug simulado")

    monkeypatch.setattr("src.pipeline.build_breakdown", boom)
    api = FakeApi([make_client(1)], [make_pl_entry(1, 1000)], total=1000)
    result = run_with(settings, api)
    assert result.exit_code is ExitCode.SUCCESS
    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED
    assert "falha interna" in result.audit_error
    assert published_files(settings) == [settings.output_dir / EXPECTED_FILE]


def test_spreadsheet_is_published_before_the_financial_audit(settings: Settings) -> None:
    seen: list[bool] = []

    class OrderCheckingApi(FakeApi):
        def fetch_consolidated_pl(self):  # type: ignore[override]
            seen.append((settings.output_dir / EXPECTED_FILE).exists())
            return super().fetch_consolidated_pl()

    run_with(settings, OrderCheckingApi([make_client(1)], [make_pl_entry(1, 1000)], total=1000))
    # Read twice (start and end of the financial collection), both after publication.
    assert seen == [True, True]


# --- Active filter and cadastral audit ----------------------------------------


def nine_cell_base() -> list[dict]:
    return [
        make_client(1, name="PF Ativa", active=True),
        make_client(2, name="PF Inativa", active=False),
        make_client(3, name="PF Status Ausente", active=None),
        make_client(4, name="PJ Ativa", cpfCnpj=make_cnpj(4), active=True),
        make_client(5, name="PJ Inativa", cpfCnpj=make_cnpj(5), active=False),
        make_client(6, name="Doc Pendente Ativo", cpfCnpj=None, active=True),
        make_client(7, name="Doc Pendente Inativo", cpfCnpj="999", active=False),
        make_client(8, name="PF Status Texto", active="true"),
    ]


def test_excel_contains_only_active_individuals(settings: Settings) -> None:
    clients = nine_cell_base()
    result = run_with(settings, FakeApi(clients, [make_pl_entry(i, 10) for i in (1, 4, 6)], total=30))

    names = [row[0] for row in sheet_rows(result.output_path)]
    assert names == ["PF Ativa"]
    assert (
        result.clients_exported
        == result.classification.count(PersonType.INDIVIDUAL, ActivityStatus.ACTIVE)
        == 1
    )
    assert result.classification.ids(status=ActivityStatus.UNKNOWN) == [3, 8]


def test_cadastral_audit_reports_id_differences(settings: Settings) -> None:
    clients = [make_client(1), make_client(2), make_client(3, active=False)]
    # Same count (2 active, 2 consolidated) but client 2 is replaced by 3.
    result = run_with(settings, FakeApi(clients, [make_pl_entry(1, 10), make_pl_entry(3, 10)], total=20))
    audit = result.cadastral
    assert audit.status is CadastralStatus.DIVERGENT
    assert (audit.active_missing_from_patrimony, audit.patrimony_not_active) == ([2], [3])
    text = render_summary(result, MoneyUnit.REAIS)
    assert "AUDITORIA CADASTRAL: DIVERGENTE" in text
    assert "A − B (ativos fora do consolidado): 1 → Cliente Ficticio 2 (ID 2)" in text
    assert "B − A (consolidado não ativo): 1 → Cliente Ficticio 3 (ID 3)" in text
    # The cadastral audit never blocks the spreadsheet.
    assert result.exit_code is ExitCode.SUCCESS and result.clients_exported == 2


def test_cadastral_audit_approved_on_exact_match(settings: Settings) -> None:
    clients = [
        make_client(1),
        make_client(2),
        make_client(3, active=False),
        make_client(4, cpfCnpj=make_cnpj(4)),
    ]
    entries = [make_pl_entry(1, 10), make_pl_entry(2, 10), make_pl_entry(4, 10)]
    result = run_with(settings, FakeApi(clients, entries, total=30))
    assert result.cadastral.status is CadastralStatus.APPROVED
    assert result.cadastral.outside_patrimony == [(3, PersonType.INDIVIDUAL, ActivityStatus.INACTIVE)]
    text = render_summary(result, MoneyUnit.REAIS)
    assert "Cadastrados fora do consolidado: Cliente Ficticio 3 (ID 3) [PF, INATIVO]" in text
    assert "  Total                      3         1             0       4" in text


def test_cadastral_audit_not_verified_without_patrimony(settings: Settings) -> None:
    api = FakeApi([make_client(1)], [], total=0, fail_on="pl", error=IntegrationError("HTTP 503"))
    result = run_with(settings, api)
    assert result.cadastral.status is CadastralStatus.NOT_VERIFIED
    assert result.clients_exported == 1


# --- Registros especiais -------------------------------------------------------


def special_settings(settings: Settings) -> Settings:
    from dataclasses import replace

    return replace(settings, special_client_ids=frozenset({"5", "6"}))


def special_base() -> list[dict]:
    return [
        make_client(1, name="PF Ativa"),
        make_client(2, name="PF Inativa", active=False),
        make_client(3, name="PJ Ativa", cpfCnpj=make_cnpj(3)),
        make_client(5, name="Familia Ficticia", cpfCnpj=None),
        make_client(6, name="Simulacao Ficticia", cpfCnpj=None),
    ]


def test_special_records_are_excluded_from_excel_but_counted(settings: Settings) -> None:
    entries = [make_pl_entry(i, 10) for i in (1, 3, 5, 6)]
    result = run_with(special_settings(settings), FakeApi(special_base(), entries, total=40))

    assert [row[0] for row in sheet_rows(result.output_path)] == ["PF Ativa"]
    assert result.special_ids == (5, 6)
    assert result.unclassified_ids == ()
    table = result.classification
    assert table.ids(person=PersonType.SPECIAL) == [5, 6]
    assert sum(table.count(person=p) for p in PersonType) == table.total == 5
    assert result.cadastral.status is CadastralStatus.APPROVED

    text = render_summary(result, MoneyUnit.REAIS)
    assert "Registro especial" in text
    assert (
        "Registros especiais configurados, fora da planilha: "
        "Familia Ficticia (ID 5), Simulacao Ficticia (ID 6)" in text
    )
    assert "Documentos sem classificação segura" not in text


def test_special_records_without_configuration_stay_pending(settings: Settings) -> None:
    entries = [make_pl_entry(i, 10) for i in (1, 3, 5, 6)]
    result = run_with(settings, FakeApi(special_base(), entries, total=40))
    assert result.unclassified_ids == (5, 6)
    assert result.cadastral.status is CadastralStatus.DIVERGENT
    assert [row[0] for row in sheet_rows(result.output_path)] == ["PF Ativa"]


def test_stale_special_configuration_is_logged(settings: Settings, caplog: pytest.LogCaptureFixture) -> None:
    from dataclasses import replace

    custom = replace(settings, special_client_ids=frozenset({"1", "999"}))
    with caplog.at_level("WARNING"):
        result = run_with(custom, FakeApi(special_base(), [], total=0))
    # Configuration wins (explicit), but stale entries are flagged by id only.
    assert result.special_ids == (1,)
    assert "999" in caplog.text
    assert "id=1 configurado como especial possui documento válido" in caplog.text


# --- Verificação independente com PTAX -------------------------------------------


def fx_settings(settings: Settings) -> Settings:
    from dataclasses import replace

    return replace(settings, fx_check_enabled=True)


def fake_ptax(monkeypatch: pytest.MonkeyPatch, rate: str = "5.3346", error: Exception | None = None) -> list:
    from datetime import date as _date

    from src.ptax import PtaxQuote, PtaxSeries

    calls: list = []

    def fetch(base_url, start, end, timeout=30.0, max_retries=2, session=None):
        calls.append((start, end))
        if error is not None:
            raise error
        quote = PtaxQuote(datetime(2026, 10, 27, 13, 5), Decimal(rate))
        return PtaxSeries((quote,), start, end, datetime(2026, 10, 28, 15, 0))

    monkeypatch.setattr("src.pipeline.fetch_usd_brl_period", fetch)
    monkeypatch.setattr(
        "src.pipeline.date", type("FixedDate", (_date,), {"today": staticmethod(lambda: _date(2026, 10, 28))})
    )
    return calls


def fx_api() -> FakeApi:
    from tests.conftest import make_wallet_detail

    entries = [make_pl_entry(1, 1000), make_pl_entry(2, 100)]
    entries[1]["wallets"][0]["walletInfoId"] = 2000
    clients = [make_client(1, currentValue=1000), make_client(2, currentValue=Decimal("123456.78"))]
    detail = make_wallet_detail(2000, walletCoin="DOLLAR", lastConsolidation="2026-10-28T04:16:00")
    return FakeApi(clients, entries, total=Decimal("1533.46"), wallet_details={2000: detail})


def test_fx_check_runs_after_publication_and_writes_local_report(
    settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    calls = fake_ptax(monkeypatch)
    result = run_with(fx_settings(settings), fx_api())

    assert result.exit_code is ExitCode.SUCCESS
    assert calls and calls[0][1] == datetime(2026, 10, 28).date()
    fx = result.fx
    assert len(fx.estimated) == 2
    assert fx.estimated_total == Decimal("1533.4600")
    assert fx.global_difference == Decimal("0.0000")
    # Cliente 2: 100 USD x 5,3346 = 533,46 contra 123.456,78 operacional → relevante.
    assert [c.client_id for c in fx.relevant] == [2]
    # A auditoria exata continua NÃO VERIFICADA (há DOLLAR sem cotação oficial).
    assert result.reconciliation.status is AuditStatus.NOT_VERIFIED
    assert result.diagnostic_path.parent == settings.output_dir / "diagnostico"
    # O Excel operacional não muda: patrimônio exibido continua sendo currentValue.
    rows = {r[0]: r for r in sheet_rows(result.output_path)}
    assert rows["Cliente Ficticio 2"][2] == pytest.approx(123456.78)
    assert len(next(iter(rows.values()))) == 8

    text = render_summary(result, MoneyUnit.REAIS, settings.output_dir.parent)
    assert "VERIFICAÇÃO INDEPENDENTE (PTAX/BCB" in text
    assert "PTAX venda 27/10/2026 13:05: 5,3346 (usada por 1 cliente(s))" in text
    assert "Divergência relevante (|dif.| ≥ 5% do operacional): 1 cliente(s)" in text
    assert "output/diagnostico/Diagnostico_Patrimonial_2026-10-28_140000.xlsx" in text


def test_fx_failure_does_not_block_excel(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    from src.ptax import PtaxError

    fake_ptax(monkeypatch, error=PtaxError("PTAX: HTTP 503."))
    result = run_with(fx_settings(settings), fx_api())
    assert result.exit_code is ExitCode.SUCCESS
    assert result.fx is None and result.fx_error == "PTAX: HTTP 503."
    assert result.diagnostic_path is None
    assert published_files(settings) == [settings.output_dir / EXPECTED_FILE]
    assert "Não executada: PTAX: HTTP 503." in render_summary(result, MoneyUnit.REAIS)


def test_fx_internal_error_does_not_block_excel(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    fake_ptax(monkeypatch)

    def boom(*_: object, **__: object) -> None:
        raise RuntimeError("bug simulado")

    monkeypatch.setattr("src.pipeline.estimate_patrimony", boom)
    result = run_with(fx_settings(settings), fx_api())
    assert result.exit_code is ExitCode.SUCCESS
    assert "falha interna" in result.fx_error
    assert (settings.output_dir / EXPECTED_FILE).exists()


def test_fx_without_wallet_details_is_not_run(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    from dataclasses import replace

    calls = fake_ptax(monkeypatch)
    result = run_with(replace(fx_settings(settings), wallet_details_enabled=False), fx_api())
    assert calls == [] and result.fx is None
    assert "moedas desconhecidas" in result.fx_error


def test_fx_logs_contain_no_personal_data(
    settings: Settings, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    fake_ptax(monkeypatch)
    with caplog.at_level("INFO"):
        run_with(fx_settings(settings), fx_api())
    log = caplog.text
    assert "Verificação PTAX" in log
    for secret in ("Cliente Ficticio", "123456.78", "123.456,78", "533.46", "533,46"):
        assert secret not in log
    from tests.conftest import make_cpf

    assert make_cpf(2) not in log


def test_diagnostic_report_can_be_disabled(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    from dataclasses import replace

    fake_ptax(monkeypatch)
    result = run_with(replace(fx_settings(settings), diagnostic_report_enabled=False), fx_api())
    assert result.fx is not None and result.diagnostic_path is None
    assert published_files(settings) == [settings.output_dir / EXPECTED_FILE]


# --- Homologação: estabilidade, duplicidade, PL zero e resumo nominal -----------


class ChangingApi(FakeApi):
    """Consolidado diferente na segunda leitura (reconsolidação durante a coleta)."""

    def __init__(self, *args, second_entries: list, second_total, **kwargs) -> None:
        super().__init__(*args, **kwargs)
        self.second = (second_entries, Decimal(second_total))
        self.pl_calls = 0

    def fetch_consolidated_pl(self):
        from src.api_client import ConsolidatedPL

        self.pl_calls += 1
        if self.pl_calls == 1:
            return super().fetch_consolidated_pl()
        return ConsolidatedPL(self.second[1], self.second[0])


def test_changes_during_collection_are_flagged_without_blocking(settings: Settings) -> None:
    clients = [make_client(1, name="Ana Ficticia"), make_client(2, name="Bruno Ficticio")]
    first = [make_pl_entry(1, 1000), make_pl_entry(2, 500)]
    second = [make_pl_entry(1, 1100), make_pl_entry(2, 500)]
    api = ChangingApi(clients, first, total=1500, second_entries=second, second_total="1500")
    result = run_with(settings, api)

    assert result.exit_code is ExitCode.SUCCESS
    assert published_files(settings) == [settings.output_dir / EXPECTED_FILE]
    assert not result.stability.stable and not result.stability.total_changed
    assert result.stability.affected_clients == [1]
    text = render_summary(result, MoneyUnit.REAIS)
    assert "DADOS ATUALIZADOS DURANTE A EXECUÇÃO: 1 cliente(s) afetado(s)" in text
    assert "Afetados: Ana Ficticia (ID 1)" in text
    assert "total da empresa inalterado" in text
    assert "Verificação temporal: ALTERADA (1 cliente(s))" in text


def test_stable_collection(settings: Settings) -> None:
    result = run_with(settings, FakeApi([make_client(1)], [make_pl_entry(1, 1000)], total=1000))
    assert result.stability.stable
    assert "ESTÁVEL: mesmos clientes, carteiras e valores" in render_summary(result, MoneyUnit.REAIS)


def test_stability_not_verified_when_second_read_fails(settings: Settings) -> None:
    class SecondReadFails(FakeApi):
        pl_reads = 0

        def fetch_consolidated_pl(self):
            self.pl_reads += 1
            if self.pl_reads == 2:
                raise IntegrationError("GET consolidated-pl: HTTP 503 inesperado.")
            return super().fetch_consolidated_pl()

    result = run_with(settings, SecondReadFails([make_client(1)], [make_pl_entry(1, 1000)], total=1000))
    assert result.stability is None and "503" in result.stability_error
    assert "NÃO VERIFICADA: GET consolidated-pl: HTTP 503" in render_summary(result, MoneyUnit.REAIS)
    assert (settings.output_dir / EXPECTED_FILE).exists()


def _wallet(wid: int, current: str, invested: str) -> dict:
    return {"walletInfoId": wid, "currentValue": Decimal(current), "investedValue": Decimal(invested)}


def family_scenario() -> tuple[list, list]:
    clients = [
        make_client(100, name="Familia Ficticia", cpfCnpj=None),
        make_client(200, name="Membro Ficticio A", currentValue=Decimal("1234.56")),
        make_client(300, name="Membro Ficticio B", currentValue=Decimal("2345.67")),
    ]
    entries = [
        {
            "clientId": 100,
            "wallets": [
                _wallet(1, "1234.56", "1100.00"),
                _wallet(2, "2345.67", "2000.00"),
                _wallet(3, "9876.54", "9000.00"),
            ],
        },
        {"clientId": 200, "wallets": [_wallet(21, "1234.56", "1100.00")]},
        {"clientId": 300, "wallets": [_wallet(31, "2345.67", "2000.00")]},
    ]
    return clients, entries


def test_family_duplication_is_reported_by_name_and_not_subtracted(settings: Settings) -> None:
    from dataclasses import replace

    clients, entries = family_scenario()
    total = Decimal("17037.00")
    custom = replace(settings, special_client_ids=frozenset({"100"}))
    result = run_with(custom, FakeApi(clients, entries, total=total))

    assert result.reconciliation.company_total == total  # nada subtraído
    assert {(e.client_a, e.client_b) for e in result.duplicates} == {(100, 200), (100, 300)}
    assert [r[0] for r in sheet_rows(result.output_path)] == ["Membro Ficticio A", "Membro Ficticio B"]
    assert result.classification.ids(person=PersonType.SPECIAL) == [100]
    text = render_summary(result, MoneyUnit.REAIS)
    assert "Familia Ficticia (ID 100): POSSÍVEL DUPLICIDADE — CONFERÊNCIA MANUAL" in text
    assert "valores idênticos aos de: Membro Ficticio A (ID 200), Membro Ficticio B (ID 300)" in text
    assert "R$ 3.580,23" in text and "NÃO subtraído do total" in text
    assert "Duplicidade em registros especiais: Familia Ficticia (ID 100): POSSÍVEL DUPLICIDADE" in text


def test_duplicate_flag_in_diagnostic_report(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    from dataclasses import replace

    fake_ptax(monkeypatch)
    clients, entries = family_scenario()
    custom = replace(fx_settings(settings), special_client_ids=frozenset({"100"}))
    result = run_with(custom, FakeApi(clients, entries, total=Decimal("17037.00")))
    sheet = load_workbook(result.diagnostic_path)["Diagnostico"]
    rows = {r[0]: r for r in sheet.iter_rows(min_row=2, values_only=True)}
    assert "Membro Ficticio A (ID 200) carteira 21" in rows[100][-1]
    assert "Familia Ficticia (ID 100) carteira 1" in rows[200][-1]
    assert rows[200][-2] == "Não"


def test_zero_patrimony_listed_by_name(settings: Settings) -> None:
    from tests.conftest import make_wallet_detail

    clients = [
        make_client(1, name="Zero Consolidado", currentValue=0),
        make_client(2, name="Zero Pendente", currentValue=0),
    ]
    entries = [make_pl_entry(1, 0), make_pl_entry(2, 0)]
    entries[1]["wallets"][0]["walletInfoId"] = 2000
    pending = {2000: make_wallet_detail(2000, lastConsolidation=None)}
    result = run_with(settings, FakeApi(clients, entries, total=0, wallet_details=pending))

    assert {z.client_id: z.category.value for z in result.zero_patrimony} == {
        1: "PL zero com carteira consolidada",
        2: "PL zero sem consolidação",
    }
    text = render_summary(result, MoneyUnit.REAIS)
    assert "PF ativas com PL zero: 2" in text
    assert (
        "Zero Pendente (ID 2): PL zero sem consolidação | 1 carteira(s), moeda REAIS,"
        " última consolidação nunca" in text
    )
    assert "Zero Consolidado (ID 1): PL zero com carteira consolidada" in text
    assert result.never_consolidated_wallets == 1


def test_relevant_divergence_listed_by_name(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    fake_ptax(monkeypatch)
    clients = [
        make_client(1, name="Cliente Divergente", currentValue=1000),
        make_client(2, name="Ok Ficticio", currentValue=500),
    ]
    entries = [make_pl_entry(1, 920), make_pl_entry(2, 500)]
    entries[1]["wallets"][0]["walletInfoId"] = 2000
    result = run_with(fx_settings(settings), FakeApi(clients, entries, total=1420))
    text = render_summary(result, MoneyUnit.REAIS)
    assert "Clientes com divergência > 1%: 1 | > 5%: 1" in text
    assert "Divergência > 5% (estimado PTAX × clients.currentValue): 1" in text
    assert "- Cliente Divergente (ID 1): -8,00%" in text
    assert "Ok Ficticio (ID 2):" not in text.split("CLIENTES QUE PRECISAM DE CONFERÊNCIA")[1]


def test_no_bare_ids_in_attention_list(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    import re

    fake_ptax(monkeypatch)
    clients, entries = family_scenario()
    clients.append(make_client(400, name="Sem Documento", cpfCnpj=None))
    result = run_with(fx_settings(settings), FakeApi(clients, entries, total=Decimal("17037.00")))
    attention = render_summary(result, MoneyUnit.REAIS).split("CLIENTES QUE PRECISAM DE CONFERÊNCIA")[1]
    items = [line for line in attention.splitlines() if line.strip().startswith("- ")]
    assert items
    for line in items:
        assert re.search(r"- .+ \(ID \d+\)", line), line


def test_cadastral_rules_preserved_with_eight_columns(settings: Settings) -> None:
    from dataclasses import replace

    clients = [*nine_cell_base(), make_client(9, name="Especial Ficticio", cpfCnpj=None)]
    entries = [make_pl_entry(i, 10) for i in (1, 4, 6, 9)]
    custom = replace(settings, special_client_ids=frozenset({"9"}))
    result = run_with(custom, FakeApi(clients, entries, total=40))
    rows = sheet_rows(result.output_path)
    assert [r[0] for r in rows] == ["PF Ativa"]
    assert (rows[0][4], rows[0][5]) == ("Cidade Exemplo", "SP")
    assert len(rows[0]) == 8

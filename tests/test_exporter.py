"""Excel layout, types, validation and atomic publication."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.config import MoneyUnit
from src.errors import ExportError, TransformationError
from src.exporter import (
    DATE_FORMAT,
    HEADERS,
    MONEY_FORMAT,
    SHEET_NAME,
    create_temp_path,
    publish,
    report_filename,
    write_workbook,
)
from src.transformer import Pendency, ReportRow
from src.validator import validate_excel, validate_row_quantities
from tests.conftest import make_client


def sample_rows() -> list[ReportRow]:
    return [
        ReportRow(1, "Ana Ficticia", date(1982, 1, 5), Decimal("1234.56"), "Rua A, 1", "Consultor X", []),
        ReportRow(2, "Bruno Ficticio", None, None, None, None, [Pendency.BIRTHDATE, Pendency.PATRIMONY]),
        ReportRow(3, "Carla Ficticia", date(1971, 2, 3), Decimal(0), "Rua C", "Consultor X", []),
    ]


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    path = tmp_path / "test.xlsx"
    write_workbook(sample_rows(), path, MoneyUnit.REAIS)
    return path


def test_filename_pattern() -> None:
    moment = datetime(2026, 10, 28, 14, 0, 0)
    assert report_filename(moment) == "Aniversariantes_2026-10-28_140000.xlsx"


def test_single_sheet_with_headers_filter_and_freeze(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames == [SHEET_NAME]
    sheet = workbook[SHEET_NAME]
    assert tuple(c.value for c in sheet[1]) == HEADERS
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:H4"
    assert all(sheet.column_dimensions[col].width >= 12 for col in "ABCDEFGH")


def test_real_date_and_numeric_money_cells(workbook_path: Path) -> None:
    sheet = load_workbook(workbook_path)[SHEET_NAME]
    birth = sheet.cell(2, 2)
    assert isinstance(birth.value, datetime) and birth.value.date() == date(1982, 1, 5)
    assert birth.number_format == DATE_FORMAT
    money = sheet.cell(2, 3)
    assert money.value == pytest.approx(1234.56)
    assert money.number_format == MONEY_FORMAT
    assert sheet.cell(4, 3).value == 0


def test_pending_cells_not_invented_and_highlighted(workbook_path: Path) -> None:
    sheet = load_workbook(workbook_path)[SHEET_NAME]
    assert sheet.cell(3, 2).value == "PENDENTE"
    assert sheet.cell(3, 3).value is None
    assert sheet.cell(3, 8).value == "PENDÊNCIA: DATA DE NASCIMENTO; PATRIMÔNIO"
    assert sheet.cell(3, 8).fill.fgColor.rgb.endswith("FFF2CC")
    assert sheet.cell(2, 8).fill.fill_type is None


def test_centavos_are_converted(tmp_path: Path) -> None:
    path = tmp_path / "c.xlsx"
    rows = [ReportRow(1, "A", date(1980, 1, 1), Decimal(123456), "R", "C", [])]
    write_workbook(rows, path, MoneyUnit.CENTAVOS)
    assert load_workbook(path)[SHEET_NAME].cell(2, 3).value == pytest.approx(1234.56)
    validate_excel(path, rows, MoneyUnit.CENTAVOS)


def test_high_precision_api_values_survive_validation(tmp_path: Path) -> None:
    # Production sends values with ~10 decimal places (beyond double precision).
    path = tmp_path / "p.xlsx"
    rows = [ReportRow(1, "A", date(1980, 1, 1), Decimal("123456789.1234567891"), "R", "C", [])]
    write_workbook(rows, path, MoneyUnit.REAIS)
    validate_excel(path, rows, MoneyUnit.REAIS)


@pytest.mark.parametrize(
    "raw",
    ["122232.41607835", "129605.15916098", "4400000.2499999999", "0.0000000001", "99999999.9999999999"],
)
def test_ten_decimal_current_values_survive_validation(tmp_path: Path, raw: str) -> None:
    # clients.currentValue arrives with up to 10 decimal places.
    path = tmp_path / "ten.xlsx"
    rows = [ReportRow(1, "A", date(1980, 1, 1), Decimal(raw), "R", "C", [])]
    write_workbook(rows, path, MoneyUnit.REAIS)
    validate_excel(path, rows, MoneyUnit.REAIS)


def test_validation_passes_for_written_file(workbook_path: Path) -> None:
    validate_excel(workbook_path, sample_rows(), MoneyUnit.REAIS)


def test_validation_detects_missing_row(workbook_path: Path) -> None:
    rows = [*sample_rows(), ReportRow(4, "Extra", None, None, None, None, [Pendency.BIRTHDATE])]
    with pytest.raises(ExportError, match="linha"):
        validate_excel(workbook_path, rows, MoneyUnit.REAIS)


def test_validation_detects_tampered_value(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    workbook[SHEET_NAME].cell(2, 3).value = 999
    workbook.save(workbook_path)
    with pytest.raises(ExportError, match="Patrimônio divergente"):
        validate_excel(workbook_path, sample_rows(), MoneyUnit.REAIS)


def test_validation_detects_text_date(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    workbook[SHEET_NAME].cell(2, 2).value = "05/01/1982"
    workbook.save(workbook_path)
    with pytest.raises(ExportError, match="Data"):
        validate_excel(workbook_path, sample_rows(), MoneyUnit.REAIS)


def test_validation_rejects_corrupt_file(tmp_path: Path) -> None:
    path = tmp_path / "corrupt.xlsx"
    path.write_bytes(b"not a zip")
    with pytest.raises(ExportError, match="reaberto"):
        validate_excel(path, [], MoneyUnit.REAIS)


def test_publish_is_atomic_move_and_never_overwrites(tmp_path: Path) -> None:
    temp = create_temp_path(tmp_path)
    temp.write_bytes(b"x")
    final = tmp_path / "out" / "final.xlsx"
    assert publish(temp, final) == final
    assert final.read_bytes() == b"x" and not temp.exists()
    other = create_temp_path(tmp_path)
    with pytest.raises(ExportError, match="já existe"):
        publish(other, final)


def test_row_quantity_validation() -> None:
    clients = [make_client(1), make_client(2)]
    rows = sample_rows()[:2]
    validate_row_quantities(clients, rows)
    with pytest.raises(TransformationError):
        validate_row_quantities(clients, rows[:1])
    with pytest.raises(TransformationError):
        validate_row_quantities([make_client(1), make_client(9)], rows)


def test_eight_columns_with_city_and_state(tmp_path: Path) -> None:
    path = tmp_path / "c.xlsx"
    rows = [
        ReportRow(
            1,
            "Ana",
            date(1980, 1, 1),
            Decimal(1),
            "Rua A, Porto Alegre, RS",
            "C",
            [],
            city="Porto Alegre",
            state="RS",
        ),
        ReportRow(2, "Bia", None, None, None, None, [Pendency.BIRTHDATE], city=None, state=None),
    ]
    write_workbook(rows, path, MoneyUnit.REAIS)
    sheet = load_workbook(path)[SHEET_NAME]
    assert tuple(c.value for c in sheet[1]) == (
        "Nome completo",
        "Data de nascimento",
        "Patrimônio",
        "Endereço completo",
        "Cidade",
        "Estado",
        "Consultor responsável",
        "Status",
    )
    assert sheet.auto_filter.ref == "A1:H3"
    assert (sheet.cell(2, 5).value, sheet.cell(2, 6).value) == ("Porto Alegre", "RS")
    assert (sheet.cell(3, 5).value, sheet.cell(3, 6).value) == (None, None)  # nada inventado
    assert sheet.cell(3, 8).fill.fgColor.rgb.endswith("FFF2CC")
    validate_excel(path, rows, MoneyUnit.REAIS)


def test_validation_detects_wrong_city_and_partial_filter(tmp_path: Path) -> None:
    rows = [ReportRow(1, "Ana", date(1980, 1, 1), Decimal(1), "R", "C", [], city="Porto Alegre", state="RS")]
    path = tmp_path / "v.xlsx"
    write_workbook(rows, path, MoneyUnit.REAIS)
    workbook = load_workbook(path)
    workbook[SHEET_NAME].cell(2, 5).value = "Canoas"
    workbook.save(path)
    with pytest.raises(ExportError, match="Cidade/estado"):
        validate_excel(path, rows, MoneyUnit.REAIS)
    write_workbook(rows, path, MoneyUnit.REAIS)
    workbook = load_workbook(path)
    workbook[SHEET_NAME].auto_filter.ref = "A1:F2"
    workbook.save(path)
    with pytest.raises(ExportError, match="Autofiltro"):
        validate_excel(path, rows, MoneyUnit.REAIS)

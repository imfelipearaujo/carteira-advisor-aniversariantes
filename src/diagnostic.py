"""Local per-client diagnostic workbook (``output/diagnostico``).

Contains personal data: it is written only under ``output/`` (ignored by Git),
atomically and re-validated, like the operational spreadsheet. The estimation
logic itself lives in :mod:`src.estimate` and does not depend on Excel.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.errors import ExportError
from src.estimate import FxEstimate
from src.exporter import DATE_FORMAT, MONEY_FORMAT, create_temp_path, discard, publish
from src.transformer import id_key

DIAGNOSTIC_DIRNAME = "diagnostico"
DIAGNOSTIC_SHEET = "Diagnostico"

DIAGNOSTIC_HEADERS = (
    "ID do cliente",
    "Nome",
    "Tipo",
    "Situação",
    "Patrimônio operacional (R$)",
    "Carteiras em REAIS (R$)",
    "Carteiras em DOLLAR (US$)",
    "PTAX venda",
    "Data PTAX",
    "Patrimônio estimado (R$)",
    "Diferença (R$)",
    "Diferença (%)",
    "Faixa",
    "Observação",
    "Última consolidação",
    "Carteiras nunca consolidadas",
    "Alterado durante a coleta",
    "Possível duplicidade",
)


def write_diagnostic_report(
    result: FxEstimate,
    clients: list[dict[str, Any]],
    labels: dict[str, tuple[str, str]],
    output_dir: Path,
    moment: datetime,
    annotations: dict[str, dict[str, Any]] | None = None,
) -> Path:
    """Write ``output/diagnostico/Diagnostico_Patrimonial_<ts>.xlsx`` atomically.

    Args:
        result: Coherence check result.
        clients: Registered clients (for names).
        labels: Client id key -> (type label, status label).
        output_dir: Operational output directory.
        moment: Run timestamp.
        annotations: Client id key -> ``consolidation`` (datetime), ``never``
            (int), ``changed`` (bool) and ``duplicate`` (str) indicators.

    Returns:
        The published path.

    Raises:
        ExportError: If writing or validation fails.
    """
    names = {id_key(c.get("id")): str(c.get("name") or "").strip() for c in clients}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DIAGNOSTIC_SHEET
    sheet.append(DIAGNOSTIC_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="1F3864")
    for c in sorted(
        result.clients, key=lambda e: -(abs(e.difference_pct) if e.difference_pct is not None else -1)
    ):
        key = id_key(c.client_id)
        person, status = labels.get(key, ("?", "?"))
        quote = c.quotes[0] if len(c.quotes) == 1 else None
        note = c.reason or ("várias datas PTAX" if len(c.quotes) > 1 else "")
        if c.crossed_thresholds:
            note = "; ".join(
                filter(
                    None, [note, "cruza limite(s): " + ", ".join(f"{t:,.0f}" for t in c.crossed_thresholds)]
                )
            )
        sheet.append(
            (
                c.client_id,
                names.get(key, ""),
                person,
                status,
                _num(c.operational_brl),
                _num(c.brl_total),
                _num(c.usd_total),
                _num(quote.sell_rate) if quote else None,
                quote.published_at.date() if quote else None,
                _num(c.estimated_brl),
                _num(c.difference),
                _num(c.difference_pct / 100) if c.difference_pct is not None else None,
                c.band,
                note,
                (annotations or {}).get(key, {}).get("consolidation"),
                (annotations or {}).get(key, {}).get("never"),
                "Sim" if (annotations or {}).get(key, {}).get("changed") else "Não",
                (annotations or {}).get(key, {}).get("duplicate") or None,
            )
        )
        row = sheet.max_row
        sheet.cell(row, 15).number_format = "DD/MM/YYYY HH:MM"
        for column in (5, 6, 10, 11):
            sheet.cell(row, column).number_format = MONEY_FORMAT
        sheet.cell(row, 7).number_format = '"US$" #,##0.00'
        sheet.cell(row, 8).number_format = "0.0000"
        sheet.cell(row, 9).number_format = DATE_FORMAT
        sheet.cell(row, 12).number_format = "0.00%"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(DIAGNOSTIC_HEADERS))}{sheet.max_row}"
    for index, header in enumerate(DIAGNOSTIC_HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, min(len(header) + 4, 40))
    sheet.column_dimensions["B"].width = 40

    final_path = output_dir / DIAGNOSTIC_DIRNAME / f"Diagnostico_Patrimonial_{moment:%Y-%m-%d_%H%M%S}.xlsx"
    temp_path = create_temp_path(output_dir)
    try:
        workbook.save(temp_path)
        reopened = load_workbook(temp_path)
        check = reopened[DIAGNOSTIC_SHEET] if reopened.sheetnames == [DIAGNOSTIC_SHEET] else None
        if (
            check is None
            or check.max_row - 1 != len(result.clients)
            or tuple(cell.value for cell in check[1]) != DIAGNOSTIC_HEADERS
            or sorted(str(c.value) for c in check["A"][1:])
            != sorted(str(e.client_id) for e in result.clients)
        ):
            raise ExportError("Relatório diagnóstico com conteúdo inesperado.")
        publish(temp_path, final_path)
    except OSError as exc:
        raise ExportError(f"Falha ao gravar o relatório diagnóstico: {exc.strerror}.") from exc
    finally:
        discard(temp_path)
    return final_path


def _num(value: Decimal | None) -> float | None:
    """Excel cells are doubles; Decimals are only used for the calculations."""
    return float(value) if value is not None else None

"""Quantity checks and post-write validation of the Excel file."""

from __future__ import annotations

import math
from datetime import date, datetime
from decimal import Decimal
from numbers import Number
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.config import MoneyUnit
from src.errors import ExportError, TransformationError
from src.exporter import HEADERS, PENDING_TEXT, SHEET_NAME, to_reais
from src.transformer import ReportRow, id_key

EXCEL_RELATIVE_PRECISION = 1e-15


def validate_row_quantities(clients: list[dict[str, Any]], rows: list[ReportRow]) -> None:
    """Ensure every collected client became exactly one row.

    Raises:
        TransformationError: If any client was dropped or duplicated.
    """
    if len(rows) != len(clients):
        raise TransformationError(
            f"Quantidade divergente: {len(clients)} cliente(s) coletado(s), {len(rows)} linha(s)."
        )
    expected = sorted(id_key(c.get("id")) for c in clients)
    actual = sorted(id_key(r.client_id) for r in rows)
    if expected != actual:
        raise TransformationError("Os IDs das linhas não correspondem aos clientes coletados.")


def validate_excel(path: Path, rows: list[ReportRow], money_unit: MoneyUnit) -> None:
    """Re-open the written workbook and check structure, types and content.

    Args:
        path: Workbook to check.
        rows: Rows that were supposed to be written, in order.
        money_unit: Unit used for ``row.patrimony``.

    Raises:
        ExportError: On any mismatch.
    """
    try:
        workbook = load_workbook(path)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise ExportError(f"Excel gerado não pôde ser reaberto ({type(exc).__name__}).") from exc

    if workbook.sheetnames != [SHEET_NAME]:
        raise ExportError(f"Abas inesperadas no Excel: {workbook.sheetnames}.")
    sheet = workbook[SHEET_NAME]
    header = tuple(cell.value for cell in sheet[1])
    if header != HEADERS:
        raise ExportError("Cabeçalho do Excel diferente do layout esperado.")
    if sheet.freeze_panes != "A2":
        raise ExportError("Primeira linha do Excel não está congelada.")
    expected_filter = f"A1:{get_column_letter(len(HEADERS))}{max(sheet.max_row, 1)}"
    if sheet.auto_filter.ref != expected_filter:
        raise ExportError("Autofiltro ausente ou sem cobrir todas as colunas do Excel.")
    comment = sheet.cell(1, 3).comment
    if comment is not None:
        raise ExportError("Comentário inesperado no cabeçalho de patrimônio.")
    if sheet.max_row - 1 != len(rows):
        raise ExportError(f"Excel contém {sheet.max_row - 1} linha(s); esperado {len(rows)}.")

    for offset, row in enumerate(rows, start=2):
        name, birth, money, address, city, state, consultant, status = (
            sheet.cell(offset, column).value for column in range(1, len(HEADERS) + 1)
        )
        if (name or "") != row.name or status != row.status:
            raise ExportError(f"Conteúdo divergente na linha {offset} do Excel.")
        if (address, consultant) != (row.address, row.consultant):
            raise ExportError(f"Endereço/consultor divergente na linha {offset} do Excel.")
        if (city, state) != (row.city, row.state):
            raise ExportError(f"Cidade/estado divergente na linha {offset} do Excel.")
        _check_date_cell(birth, row.birthdate, offset)
        _check_money_cell(money, row.patrimony, money_unit, offset)


def _check_date_cell(value: Any, expected: date | None, line: int) -> None:
    if expected is None:
        if value != PENDING_TEXT:
            raise ExportError(f"Data pendente não sinalizada na linha {line} do Excel.")
        return
    if not isinstance(value, datetime) or value.date() != expected:
        raise ExportError(f"Data de nascimento inválida na linha {line} do Excel.")


def _check_money_cell(value: Any, expected: Decimal | None, unit: MoneyUnit, line: int) -> None:
    if expected is None:
        if value is not None:
            raise ExportError(f"Patrimônio não apurado preenchido na linha {line} do Excel.")
        return
    if not isinstance(value, Number) or isinstance(value, bool):
        raise ExportError(f"Patrimônio não numérico na linha {line} do Excel.")
    # Excel keeps at most 15 significant digits and the API sends up to ~10
    # decimal places, so the cell is compared at Excel's own relative
    # precision (comparing rounded strings fails at rounding boundaries). This
    # only checks the written cell against its source; it is not a financial
    # tolerance, and no audit is ever computed from these cells.
    if not math.isclose(float(value), float(to_reais(expected, unit)), rel_tol=EXCEL_RELATIVE_PRECISION):
        raise ExportError(f"Patrimônio divergente na linha {line} do Excel.")

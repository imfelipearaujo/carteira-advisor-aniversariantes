"""Excel generation with atomic publication.

The workbook is always written to a temporary file in the destination
directory first and only moved into place (``os.replace``, atomic on the same
volume) after it has been validated.
"""

from __future__ import annotations

import contextlib
import os
import tempfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.config import MoneyUnit
from src.errors import ExportError
from src.transformer import ReportRow

SHEET_NAME = "Aniversariantes"
HEADERS = (
    "Nome completo",
    "Data de nascimento",
    "Patrimônio",
    "Endereço completo",
    "Cidade",
    "Estado",
    "Consultor responsável",
    "Status",
)
STATUS_COLUMN = HEADERS.index("Status") + 1
PENDING_TEXT = "PENDENTE"
DATE_FORMAT = "DD/MM/YYYY"
MONEY_FORMAT = '"R$" #,##0.00;-"R$" #,##0.00'

_HEADER_FILL = PatternFill("solid", start_color="1F3864")
_PENDING_FILL = PatternFill("solid", start_color="FFF2CC")
_MIN_WIDTH, _MAX_WIDTH = 12, 60


PROVISIONAL_NOTE = (
    "Patrimônio PROVISÓRIO: valor de clients.currentValue do cadastro do Carteira Advisor. "
    "Não auditado nem conciliado com o patrimônio consolidado da empresa."
)


def report_filename(moment: datetime) -> str:
    """Return ``Aniversariantes_AAAA-MM-DD_HHMMSS.xlsx``."""
    return f"Aniversariantes_{moment:%Y-%m-%d_%H%M%S}.xlsx"


def to_reais(value: Decimal, unit: MoneyUnit) -> Decimal:
    """Convert an API monetary value to reais for display."""
    return value / 100 if unit is MoneyUnit.CENTAVOS else value


def write_workbook(rows: list[ReportRow], path: Path, money_unit: MoneyUnit) -> None:
    """Write the single-sheet birthday workbook.

    Args:
        rows: Already sorted rows.
        path: Destination file (typically a temporary path).
        money_unit: Unit of ``row.patrimony``.

    Raises:
        ExportError: If the file cannot be written.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    # The column name is part of the fixed layout; the provisional nature of the
    # values is flagged on the header cell and in the file properties.
    sheet.cell(1, 3).comment = Comment(PROVISIONAL_NOTE, "Carteira Advisor")
    workbook.properties.title = "Aniversariantes — Pessoa Física"
    workbook.properties.description = PROVISIONAL_NOTE

    for row in rows:
        # Excel cells are IEEE doubles; the exact Decimal is only used for the audit.
        patrimony = float(to_reais(row.patrimony, money_unit)) if row.patrimony is not None else None
        sheet.append(
            (
                row.name or None,
                row.birthdate if row.birthdate is not None else PENDING_TEXT,
                patrimony,
                row.address,
                row.city,
                row.state,
                row.consultant,
                row.status,
            )
        )
        excel_row = sheet.max_row
        date_cell = sheet.cell(excel_row, 2)
        if row.birthdate is not None:
            date_cell.number_format = DATE_FORMAT
        sheet.cell(excel_row, 3).number_format = MONEY_FORMAT
        if row.pendencies:
            sheet.cell(excel_row, STATUS_COLUMN).fill = _PENDING_FILL

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(sheet.max_row, 1)}"
    _fit_columns(sheet, rows, money_unit)

    try:
        workbook.save(path)
    except OSError as exc:
        raise ExportError(f"Não foi possível gravar o Excel temporário: {exc.strerror}.") from exc


def _fit_columns(sheet, rows: list[ReportRow], money_unit: MoneyUnit) -> None:
    widths = [len(header) for header in HEADERS]
    for row in rows:
        money = to_reais(row.patrimony, money_unit) if row.patrimony is not None else None
        texts = (
            row.name,
            "00/00/0000",
            f"R$ {money:,.2f}" if money is not None else "",
            row.address or "",
            row.city or "",
            row.state or "",
            row.consultant or "",
            row.status,
        )
        widths = [max(width, len(text)) for width, text in zip(widths, texts, strict=True)]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, _MIN_WIDTH), _MAX_WIDTH)


def create_temp_path(directory: Path) -> Path:
    """Create an empty temporary ``.xlsx`` path in ``directory``.

    The temporary file lives in the destination directory so the final
    ``os.replace`` stays on the same volume and is atomic.
    """
    try:
        directory.mkdir(parents=True, exist_ok=True)
        handle, name = tempfile.mkstemp(prefix=".tmp_aniversariantes_", suffix=".xlsx", dir=directory)
    except OSError as exc:
        raise ExportError(f"Não foi possível preparar o diretório de saída: {exc.strerror}.") from exc
    os.close(handle)
    return Path(name)


def publish(temp_path: Path, final_path: Path) -> Path:
    """Atomically move the validated temporary file to its final name.

    Raises:
        ExportError: If the destination already exists or the move fails.
    """
    if final_path.exists():
        raise ExportError(f"Arquivo de destino já existe: {final_path.name}.")
    try:
        final_path.parent.mkdir(parents=True, exist_ok=True)
        os.replace(temp_path, final_path)
    except OSError as exc:
        raise ExportError(f"Falha ao publicar o arquivo: {exc.strerror}.") from exc
    return final_path


def discard(temp_path: Path) -> None:
    """Remove a temporary file, ignoring absence."""
    with contextlib.suppress(OSError):
        temp_path.unlink(missing_ok=True)
